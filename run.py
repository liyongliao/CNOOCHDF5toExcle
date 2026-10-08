"""Small desktop entry point; scientific libraries load after the window appears."""
from __future__ import annotations

import argparse
import json
import logging
import os
from pathlib import Path
import struct
import sys
import time


def write_report(path, report):
    if path:
        Path(path).parent.mkdir(parents=True, exist_ok=True)
        Path(path).write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    if sys.stdout is not None:
        print(json.dumps(report, ensure_ascii=False))


def self_test():
    """Exercise the bundled HDF5 reader and both export writers offline."""
    import csv
    import tempfile
    import numpy as np
    import h5py
    from openpyxl import load_workbook
    import app as backend

    def require(condition, detail):
        if not condition:
            raise RuntimeError(detail)

    started = time.perf_counter()
    with tempfile.TemporaryDirectory(prefix="h5converter-test-") as temporary:
        directory = Path(temporary)
        source = directory / "Instruct_001-Test-2026.h5"
        timestamps = np.array([1767225600, 1767225610, 1767225620], dtype="float64")
        series = np.zeros(3, dtype=[("time", "<f8"), ("value", "<f8")])
        series["time"] = timestamps
        series["value"] = [100, 110, 120]
        with h5py.File(str(source), "w") as fixture:
            pressure = fixture.create_dataset("EQRTZ S1 PRES PSI A", data=series)
            pressure.attrs["UoM"] = "PSI"
            pressure.attrs["Measurement Type"] = "Pressure"
            series["value"] = [20, 21, 22]
            temperature = fixture.create_dataset("EQRTZ S1 TEMP CELSIUS A", data=series)
            temperature.attrs["UoM"] = "degC"
            temperature.attrs["Measurement Type"] = "Temperature"
        scan = backend.scan_directory(backend.ScanPayload(path=str(directory)))
        require(len(scan["files"]) == 1, "HDF5 scan failed")
        inspection = backend.inspect_hdf5(backend.InspectPayload(path=str(source)))
        require(len(inspection["datasets"]) == 2, "Compound dataset detection failed")
        configs = [backend.ExportConfig(
            filePath=str(source),
            selectedFields=["EQRTZ S1 PRES PSI A", "EQRTZ S1 TEMP CELSIUS A"],
            interval=10, customName="verified." + extension, tempUnit="degC", presUnit="PSI",
        ) for extension in ("csv", "xlsx")]
        task_ids = backend.trigger_export(backend.ExportPayload(
            configs=configs, outputDir=str(directory)
        ))["taskIds"]
        deadline = time.monotonic() + 90
        while time.monotonic() < deadline:
            statuses = backend.get_status(",".join(task_ids))
            if all(item["status"] in ("completed", "failed", "cancelled") for item in statuses.values()):
                break
            time.sleep(0.05)
        else:
            raise RuntimeError("Export self-test timed out")
        require(all(item["status"] == "completed" for item in statuses.values()), str(statuses))
        with (directory / "verified.csv").open(encoding="utf-8-sig", newline="") as stream:
            require("Well name:Test" in stream.readline(), "CSV metadata missing")
            rows = list(csv.reader(stream))
        require(len(rows) == 4 and rows[1][1:] == ["100.00", "20.00"], "CSV contents incorrect: " + str(rows))
        workbook = load_workbook(str(directory / "verified.xlsx"), read_only=True)
        try:
            sheet = workbook["Data"]
            require(sheet.cell(3, 2).value == 100, "Excel pressure value incorrect")
            require(sheet.cell(3, 2).number_format == "0.00", "Excel number format incorrect")
            require(sheet.cell(3, 3).value == 20, "Excel temperature value incorrect")
        finally:
            workbook.close()
        return {"status": "passed", "mode": "self-test", "architecture": struct.calcsize("P") * 8,
                "python": sys.version.split()[0], "seconds": round(time.perf_counter() - started, 3),
                "checks": ["HDF5 scan", "compound datasets", "CSV values", "Excel values and formatting"]}


def run_web():
    """Optional browser interface retained for existing users."""
    import socket
    import threading
    import urllib.request
    import webbrowser
    import app as backend
    import uvicorn

    listener = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
    listener.bind(("127.0.0.1", 0))
    listener.listen(128)
    port = listener.getsockname()[1]
    url = "http://127.0.0.1:{}".format(port)
    server = uvicorn.Server(uvicorn.Config(backend.app, host="127.0.0.1", port=port,
                                         log_level="warning", log_config=None,
                                         loop="asyncio", http="h11", ws="none"))

    def open_when_ready():
        deadline = time.monotonic() + 30
        while time.monotonic() < deadline:
            try:
                with urllib.request.urlopen(url, timeout=0.25) as response:
                    if response.status == 200:
                        webbrowser.open(url)
                        return
            except (OSError, TimeoutError):
                time.sleep(0.05)

    threading.Thread(target=open_when_ready, daemon=True).start()
    try:
        server.run(sockets=[listener])
    finally:
        listener.close()


def main(argv=None):
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description="井下压力 HDF5 转 Excel / CSV")
    parser.add_argument("--web", action="store_true", help="使用原有浏览器界面")
    parser.add_argument("--smoke-test", action="store_true", help="验证桌面窗口后自动退出")
    parser.add_argument("--self-test", action="store_true", help="验证 HDF5 和导出依赖后自动退出")
    parser.add_argument("--report", help="将验证结果写入 JSON 文件")
    arguments = parser.parse_args(argv)
    started = time.perf_counter()
    try:
        if arguments.self_test:
            write_report(arguments.report, self_test())
        elif arguments.web:
            run_web()
        else:
            import desktop
            window_report = desktop.main(smoke_test=arguments.smoke_test)
            if arguments.smoke_test:
                if not window_report or not window_report.get("success") or "firstPaintMs" not in window_report:
                    raise RuntimeError("Desktop window did not finish rendering")
                write_report(arguments.report, {
                    "status": "passed", "mode": "smoke-test",
                    "architecture": struct.calcsize("P") * 8,
                    "seconds": round(time.perf_counter() - started, 3),
                    "scientific_libraries_loaded": any(name in sys.modules for name in ("numpy", "h5py", "app")),
                    "window": window_report,
                })
        return 0
    except Exception as error:
        log_directory = Path(os.environ.get("LOCALAPPDATA", str(Path.home()))) / "H5ToExcelConverter"
        log_directory.mkdir(parents=True, exist_ok=True)
        logging.basicConfig(handlers=[logging.FileHandler(str(log_directory / "error.log"), encoding="utf-8")],
                            level=logging.ERROR, format="%(asctime)s %(levelname)s %(message)s")
        logging.exception("Application failed")
        write_report(arguments.report, {"status": "failed", "error": str(error)})
        if not (arguments.self_test or arguments.smoke_test):
            try:
                import tkinter as tk
                from tkinter import messagebox
                root = tk.Tk()
                root.withdraw()
                messagebox.showerror("启动失败", "{}\n\n详细日志：{}".format(error, log_directory / "error.log"))
                root.destroy()
            except Exception:
                if sys.stderr is not None:
                    print(str(error), file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
