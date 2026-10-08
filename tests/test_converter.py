"""Synthetic-data regression tests; production well files are never used."""
import csv
import datetime
import json
import math
import os
from pathlib import Path
import tempfile
import threading
import unittest
from unittest.mock import patch

import h5py
import numpy as np
from openpyxl import load_workbook
from fastapi import HTTPException

import app


class ConverterTests(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.root = Path(self.folder.name)
        self.source = self.root / "Instruct_123-Well-A.h5"
        self.output = self.root / "out"
        self.output.mkdir()
        self.tasks = []

    def tearDown(self):
        for task in self.tasks:
            app.cancel_task(task)
            future = app._TASK_FUTURES.get(task)
            if future and not future.cancelled():
                future.result(timeout=15)
        self.folder.cleanup()

    def create_plain(self, times=None, values=None, clock="time", data="voltage"):
        with h5py.File(str(self.source), "w") as file:
            file.create_dataset(clock, data=[0., 10., 20.] if times is None else times)
            file.create_dataset(data, data=[1.1, 2.2, 3.3] if values is None else values)

    def config(self, **kwargs):
        defaults = dict(filePath=str(self.source), selectedFields=["voltage"], interval=10,
                        customName="result.csv")
        defaults.update(kwargs)
        return app.ExportConfig(**defaults)

    def run_export(self, **kwargs):
        payload = app.ExportPayload(configs=[self.config(**kwargs)], outputDir=str(self.output))
        task_id = app.trigger_export(payload)["taskIds"][0]
        self.tasks.append(task_id)
        app._TASK_FUTURES[task_id].result(timeout=15)
        return app.get_status(task_id)[task_id]

    def csv_rows(self, name="result.csv"):
        with open(self.output / name, encoding="utf-8", newline="") as source:
            return list(csv.reader(source))

    def test_compound_time_detection_excludes_temperature_and_voltage(self):
        data = np.array([(20., 293.15), (0., 273.15), (10., 283.15)],
                        dtype=[("elapsed_seconds", "f8"), ("temperature", "f8")])
        with h5py.File(str(self.source), "w") as file:
            dataset = file.create_dataset("temperature_sensor", data=data)
            dataset.attrs["UoM"] = np.bytes_("K")
        inspected = app.inspect_hdf5(app.InspectPayload(path=str(self.source)))
        self.assertEqual(inspected["detectedTimeField"], "temperature_sensor:elapsed_seconds")
        self.assertEqual(inspected["timeMinStr"], "0.0")
        self.assertEqual(inspected["timeMaxStr"], "20.0")
        self.assertEqual(inspected["datasets"][0]["dtype"], "Compound (Time Series)")
        self.assertFalse(app.is_time_field("voltage"))
        self.assertFalse(app.is_time_field("temperature"))
        status = self.run_export(selectedFields=["temperature_sensor"])
        self.assertEqual(status["status"], "completed", status)
        rows = self.csv_rows()
        self.assertEqual(rows[1], ["Date time", "temperature_sensor (degC)"])
        self.assertEqual([row[1] for row in rows[2:]], ["0.00", "10.00", "20.00"])

    def test_compound_multiple_values_remain_selectable(self):
        values = np.array([(0., 10., 20.), (10., 30., 40.)],
                          dtype=[("time", "f8"), ("voltage", "f8"), ("current", "f8")])
        with h5py.File(str(self.source), "w") as file:
            file.create_dataset("sensor", data=values)
        inspected = app.inspect_hdf5(app.InspectPayload(path=str(self.source)))
        self.assertEqual({entry["path"] for entry in inspected["datasets"]},
                         {"sensor:time", "sensor:voltage", "sensor:current"})
        status = self.run_export(selectedFields=["sensor:voltage", "sensor:current"])
        self.assertEqual(status["status"], "completed", status)
        self.assertEqual(self.csv_rows()[2][1:], ["10.00", "20.00"])

    def test_compound_dataset_named_time_uses_its_time_member(self):
        with h5py.File(str(self.source), "w") as file:
            file.create_dataset("time", data=np.array([(10., 2.), (0., 1.)],
                                dtype=[("timestamp", "f8"), ("value", "f8")]))
        inspected = app.inspect_hdf5(app.InspectPayload(path=str(self.source)))
        self.assertEqual(inspected["detectedTimeField"], "time:timestamp")
        self.assertEqual(inspected["timeMinStr"], "0.0")

    def test_explicit_auto_clock_type_preserves_other_embedded_clock_units(self):
        seconds = np.array([1_700_000_000., 1_700_000_010., 1_700_000_020.])
        records = np.empty(3, dtype=[("timestamp", "f8"), ("value", "f8")])
        with h5py.File(str(self.source), "w") as file:
            records["timestamp"] = seconds
            records["value"] = [1., 2., 3.]
            file.create_dataset("a_sensor", data=records)
            records["timestamp"] = seconds * 1000
            records["value"] = [10., 20., 30.]
            file.create_dataset("b_sensor", data=records)
        inspected = app.inspect_hdf5(app.InspectPayload(path=str(self.source)))
        self.assertEqual(inspected["detectedTimeField"], "a_sensor:timestamp")
        status = self.run_export(selectedFields=["a_sensor", "b_sensor"], timeType="timestamp_seconds")
        self.assertEqual(status["status"], "completed", status)
        rows = self.csv_rows()
        self.assertEqual(len(rows), 5)
        self.assertEqual(rows[2][0], "2023/11/14 22:13:20")
        self.assertEqual([row[1:] for row in rows[2:]],
                         [["1.00", "10.00"], ["2.00", "20.00"], ["3.00", "30.00"]])

    def test_compound_dataset_is_read_once_per_export(self):
        with h5py.File(str(self.source), "w") as file:
            file.create_dataset("sensor", data=np.array([(0., 1.), (10., 2.)],
                                dtype=[("time", "f8"), ("value", "f8")]))
        original = h5py.Dataset.__getitem__
        raw_reads = []

        def record(dataset, selection):
            if selection is Ellipsis:
                raw_reads.append(dataset.name)
            return original(dataset, selection)

        with patch.object(h5py.Dataset, "__getitem__", record):
            status = self.run_export(selectedFields=["sensor"])
        self.assertEqual(status["status"], "completed", status)
        self.assertEqual(raw_reads, ["/sensor"])

    def test_unsorted_missing_and_duplicate_time_alignment(self):
        self.create_plain(times=[20., 0., math.nan, 10., 10., math.inf],
                          values=[3.3, 1.1, 99., 2.2, 88., 77.])
        status = self.run_export()
        self.assertEqual(status["status"], "completed", status)
        self.assertEqual([row[1] for row in self.csv_rows()[2:]], ["1.10", "2.20", "3.30"])
        actual = app.find_nearest_indices(np.array([20., math.nan, 0., 10.]), np.array([0., 5., 20.]))
        np.testing.assert_array_equal(actual, [2, 2, 0])
        np.testing.assert_array_equal(app.find_nearest_indices([math.nan], [0., 10.]), [-1, -1])

    def test_configured_global_clock_and_type_are_honoured(self):
        self.create_plain(times=[1_700_000_000_000., 1_700_000_010_000., 1_700_000_020_000.],
                          clock="clocks/arbitrary", data="data/voltage")
        with h5py.File(str(self.source), "a") as file:
            file.create_dataset("data/time", data=[0., 1., 2.])
        status = self.run_export(selectedFields=["data/voltage"], timeField="clocks/arbitrary",
                                 timeType="timestamp_ms", interval=10)
        self.assertEqual(status["status"], "completed", status)
        self.assertEqual(self.csv_rows()[2][0], "2023/11/14 22:13:20")
        self.assertEqual([row[1] for row in self.csv_rows()[2:]], ["1.10", "2.20", "3.30"])

    def test_relative_clock_base_date_and_numeric_bounds(self):
        self.create_plain()
        status = self.run_export(baseDate="2025-01-02 03:04:05", timeType="relative_seconds",
                                 startTimeStr="10", endTimeStr="20")
        self.assertEqual(status["status"], "completed", status)
        self.assertEqual(self.csv_rows()[2][0], "2025/1/2 3:04:15")
        self.assertEqual(len(self.csv_rows()), 4)

    def test_explicit_clock_with_wrong_length_fails_clearly(self):
        self.create_plain()
        with h5py.File(str(self.source), "a") as file:
            file.create_dataset("wrong_clock", data=[0., 10.])
        status = self.run_export(timeField="wrong_clock")
        self.assertEqual(status["status"], "failed", status)
        self.assertIn("长度", status["error"])

    def test_string_time_preserves_timezone_and_invalid_values(self):
        self.create_plain(times=np.array([b"2025-01-01T00:00:20Z", b"invalid",
                                         b"2025-01-01T08:00:00+08:00"], dtype="S30"),
                          values=[3.3, 99., 1.1])
        status = self.run_export()
        self.assertEqual(status["status"], "completed", status)
        rows = self.csv_rows()
        self.assertEqual(rows[2][0], "2025/1/1 0:00:00")
        self.assertEqual(rows[-1][1], "3.30")
        inspected = app.inspect_hdf5(app.InspectPayload(path=str(self.source)))
        self.assertEqual(inspected["timeMinStr"], "2025-01-01 00:00:00")
        self.assertEqual(inspected["timeMaxStr"], "2025-01-01 00:00:20")

    def test_excel_metadata_headers_unit_conversion_and_two_decimals(self):
        self.create_plain(data="pressure", values=[1000., 2000., math.nan])
        with h5py.File(str(self.source), "a") as file:
            file["pressure"].attrs["UoM"] = np.bytes_("kPa")
        status = self.run_export(selectedFields=["pressure"], customName="result", presUnit="MPa")
        self.assertEqual(status["status"], "completed", status)
        self.assertTrue(status["output_path"].endswith("result.xlsx"))
        workbook = load_workbook(status["output_path"])
        try:
            worksheet = workbook["Data"]
            self.assertEqual(worksheet["A1"].value, "Well name:Well, Sn :123 ,  Version :2.110r512")
            self.assertEqual(worksheet["B2"].value, "pressure (MPa)")
            self.assertEqual(worksheet["B3"].value, 1.)
            self.assertEqual(worksheet["B3"].number_format, "0.00")
            self.assertEqual(worksheet["A3"].value, "1970/1/1 0:00:00")
            self.assertIsNone(worksheet["B5"].value)
        finally:
            workbook.close()

    def test_duplicate_headers_do_not_drop_columns(self):
        self.create_plain(data="first/voltage")
        with h5py.File(str(self.source), "a") as file:
            file.create_dataset("second/voltage", data=[4., 5., 6.])
        status = self.run_export(selectedFields=["first/voltage", "second/voltage"], timeField="time")
        self.assertEqual(status["status"], "completed", status)
        self.assertEqual(len(self.csv_rows()[1]), 3)
        self.assertEqual(self.csv_rows()[2][1:], ["1.10", "4.00"])

    def test_excel_sheet_split_repeats_metadata_and_header(self):
        self.create_plain()
        with patch.object(app, "MAX_ROWS_PER_SHEET", 2):
            status = self.run_export(customName="result.xlsx")
        self.assertEqual(status["status"], "completed", status)
        workbook = load_workbook(status["output_path"])
        try:
            self.assertEqual(workbook.sheetnames, ["Data_Part1", "Data_Part2"])
            self.assertEqual(workbook["Data_Part2"]["A1"].value, workbook["Data_Part1"]["A1"].value)
            self.assertEqual(workbook["Data_Part2"]["B3"].value, 3.3)
        finally:
            workbook.close()

    def test_invalid_intervals_and_empty_fields_rejected_before_enqueue(self):
        self.create_plain()
        for interval in (0., -1., math.nan, math.inf):
            with self.subTest(interval=interval), self.assertRaises(HTTPException):
                app.trigger_export(app.ExportPayload(configs=[self.config(interval=interval)], outputDir=str(self.output)))
        with self.assertRaises(HTTPException):
            app.trigger_export(app.ExportPayload(configs=[self.config(selectedFields=[])], outputDir=str(self.output)))

    def test_unsafe_and_unsupported_output_names_rejected(self):
        self.create_plain()
        for name in ("../outside.csv", "sub/file.csv", "C:\\outside.csv", "CON.xlsx", "LPT1.csv", "x.xls", "x.", ""):
            with self.subTest(name=name), self.assertRaises(HTTPException):
                app.trigger_export(app.ExportPayload(configs=[self.config(customName=name)], outputDir=str(self.output)))

    def test_duplicate_batch_output_names_rejected(self):
        self.create_plain()
        with self.assertRaises(HTTPException):
            app.trigger_export(app.ExportPayload(configs=[self.config(customName="Result.csv"),
                self.config(customName="result.csv")], outputDir=str(self.output)))

    def test_grid_limit_is_checked_before_numpy_allocation(self):
        self.create_plain()
        with patch.object(app.np, "arange", side_effect=AssertionError("grid allocated before validation")):
            status = self.run_export(startTimeStr="1970-01-01", endTimeStr="2100-01-01", interval=.001)
        self.assertEqual(status["status"], "failed", status)
        self.assertIn("数据量", status["error"])
        self.assertFalse((self.output / "result.csv").exists())

    def test_wide_exports_adapt_chunks_to_memory_budget(self):
        with patch.object(app, "MAX_CHUNK_MEMORY", 16 * 1024**2):
            self.assertEqual(app._chunk_row_limit(2, 1), 8192)
            self.assertLessEqual(app._chunk_row_limit(16383, 1), 128)
            many_axes = app._chunk_row_limit(16383, 16383)
            self.assertLessEqual(many_axes, 64)
            self.assertGreaterEqual(many_axes, 1)
            self.assertLess(many_axes, app._chunk_row_limit(16383, 1))
        with patch.object(app, "MAX_CHUNK_MEMORY", 32 * 1024**2):
            self.assertLessEqual(app._chunk_row_limit(16383, 16383), 128)

    def test_reverse_range_fails_and_preserves_existing_output(self):
        self.create_plain()
        original = self.output / "result.csv"
        original.write_text("previous completed output", encoding="utf-8")
        status = self.run_export(startTimeStr="2025-01-02", endTimeStr="2025-01-01")
        self.assertEqual(status["status"], "failed", status)
        self.assertEqual(original.read_text(), "previous completed output")

    def test_cancellation_stops_writes_and_keeps_previous_excel_file(self):
        self.create_plain(times=np.arange(4000.), values=np.arange(4000.))
        previous = self.output / "result.xlsx"
        previous.write_bytes(b"previous completed file")
        gate = threading.Event()
        release = threading.Event()
        write_original = app._write_excel_rows

        def gated_write(*args, **kwargs):
            gate.set()
            release.wait(timeout=10)
            return write_original(*args, **kwargs)

        with patch.object(app, "_write_excel_rows", gated_write):
            task = app.trigger_export(app.ExportPayload(configs=[self.config(customName="result.xlsx", interval=1)],
                                                       outputDir=str(self.output)))["taskIds"][0]
            self.tasks.append(task)
            self.assertTrue(gate.wait(timeout=5))
            self.assertEqual(app.cancel_task(task), {"success": True})
            release.set()
            app._TASK_FUTURES[task].result(timeout=10)
        self.assertEqual(app.get_status(task)[task]["status"], "cancelled")
        self.assertEqual(previous.read_bytes(), b"previous completed file")
        self.assertFalse(list(self.output.glob(".hdf5-export-*")))

    def test_cancellation_mid_excel_stream_cleans_up_temporary_files(self):
        self.create_plain(times=np.arange(2000.), values=np.arange(2000.))
        original = app._format_timestamp
        seen = [0]

        def cancel_during_rows(timestamp):
            seen[0] += 1
            if seen[0] == 1000:
                with app.tasks_lock:
                    for task_id, task in app.TASKS.items():
                        if task["status"] == "running" and task["output_path"] == str(self.output / "result.xlsx"):
                            app.cancel_task(task_id)
            return original(timestamp)

        with patch.object(app, "_format_timestamp", cancel_during_rows):
            status = self.run_export(customName="result.xlsx", interval=1)
        self.assertEqual(status["status"], "cancelled", status)
        self.assertLess(seen[0], 2000)
        self.assertFalse(list(self.output.iterdir()))

    def test_inspect_reads_only_time_chunks_without_measurement_data(self):
        self.create_plain(times=np.arange(150000.) + 1_700_000_000., values=np.arange(150000.))
        original = h5py.Dataset.__getitem__
        reads = []

        def record(dataset, selection):
            reads.append((dataset.name, selection))
            return original(dataset, selection)

        with patch.object(h5py.Dataset, "__getitem__", record):
            inspected = app.inspect_hdf5(app.InspectPayload(path=str(self.source)))
        self.assertEqual(inspected["timeType"], "timestamp_seconds")
        self.assertTrue(reads)
        self.assertTrue(all(name == "/time" for name, selection in reads))
        for _, selection in reads:
            self.assertLessEqual(selection[0].stop - selection[0].start, 65536)

    def test_missing_time_does_not_silently_export_blank_values(self):
        with h5py.File(str(self.source), "w") as file:
            file.create_dataset("voltage", data=[1., 2., 3.])
        inspected = app.inspect_hdf5(app.InspectPayload(path=str(self.source)))
        self.assertIsNone(inspected["detectedTimeField"])
        self.assertEqual(inspected["timeFields"], [])
        self.assertEqual(inspected["timeType"], "none")
        self.assertIsNone(inspected["timeMinStr"])
        self.assertIsNone(inspected["timeMaxStr"])
        status = self.run_export()
        self.assertEqual(status["status"], "failed", status)
        self.assertIn("时间轴", status["error"])

    def test_atomic_configuration_fallback_and_honest_save_failure(self):
        primary = self.root / "readonly" / "config.json"
        fallback = self.root / "settings" / "config.json"
        default = self.root / "config.default.json"
        default.write_text(json.dumps({"h5_src_path": "", "h5_field_presets": {"example": {}}}))
        real_replace = os.replace

        def reject_primary(source, target):
            if str(target) == str(primary):
                raise PermissionError("read only application directory")
            return real_replace(source, target)

        with patch.object(app, "CONFIG_FILE", str(primary)), patch.object(app, "FALLBACK_CONFIG_FILE", str(fallback)), \
                patch.object(app, "BUNDLED_CONFIG_FILE", str(default)), patch.object(app.os, "replace", reject_primary):
            result = app.save_config({"h5_src_path": "chosen folder"})
            self.assertEqual(result["configPath"], str(fallback))
            self.assertEqual(app.get_config()["h5_src_path"], "chosen folder")
            self.assertIn("example", app.get_config()["h5_field_presets"])
            self.assertFalse(list(self.root.rglob(".config-*")))
        with patch.object(app, "CONFIG_FILE", str(primary)), patch.object(app, "FALLBACK_CONFIG_FILE", str(fallback)), \
                patch.object(app.os, "replace", side_effect=PermissionError("read only")):
            with self.assertRaises(HTTPException) as result:
                app.save_config({"h5_out_path": "new folder"})
            self.assertEqual(result.exception.status_code, 500)

    def test_task_history_is_bounded(self):
        original = app.TASKS.copy()
        try:
            with app.tasks_lock:
                app.TASKS.clear()
                app.TASKS.update({f"test{i}": {"status": "completed", "finished_at": i} for i in range(150)})
                app._prune_tasks()
                self.assertEqual(len(app.TASKS), app.MAX_TASK_HISTORY)
                self.assertNotIn("test0", app.TASKS)
        finally:
            with app.tasks_lock:
                app.TASKS.clear()
                app.TASKS.update(original)


if __name__ == "__main__":
    unittest.main()
